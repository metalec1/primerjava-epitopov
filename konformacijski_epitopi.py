"""
Primerjava konformacijskih epitopov dveh proteinov.

Primerjava poteka na ravni posamezne aminokisline: ujemanje je le par z
identično aminokislino na enakem položaju v proteinu (npr. K108 = K108).
Za vsak ostanek referenčnega epitopa skripta izpiše, v katerih epitopih
primerjalnega proteina se pojavi. Rezultat je Excel datoteka z enim
delovnim listom na referenčni epitop.

Uporaba:
    python konformacijski_epitopi.py <referenca.csv> <primerjalni.csv> <izhod.xlsx>
"""

import csv
import os
import re
import sys

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ZAHTEVANI_STOLPCI = ["Protein", "ID Epitopa", "Residues"]
PREPOVEDANI_ZNAKI_SHEET = [":", "/", "\\", "?", "*", "[", "]"]


# ==========================================
# 0. VARNO BRANJE CSV
# ==========================================

def preberi_csv_varno(pot):
    """
    Prebere CSV in samodejno popravi primer, ko je bila vsaka podatkovna
    vrstica ovita v dodaten par narekovajev okoli CELE vrstice
    (pogosto pri izvozu iz Excela). Pravilno oblikovan CSV prebere normalno.
    """
    with open(pot, newline="", encoding="utf-8-sig") as f:
        outer_reader = csv.reader(f)
        try:
            header = next(outer_reader)
        except StopIteration:
            raise ValueError(f"Datoteka '{pot}' je prazna.")

        rows = []
        for outer_row in outer_reader:
            if not outer_row:
                continue
            if len(outer_row) == 1 and len(header) > 1:
                rows.append(next(csv.reader([outer_row[0]])))
            else:
                rows.append(outer_row)

    return pd.DataFrame(rows, columns=header)


# ==========================================
# 1. ČIŠČENJE IN PRIMERJAVA OSTANKOV
# ==========================================

def ocisti_residues(residues_str):
    """'A:K108, A:A109, A:L110' -> ['K108', 'A109', 'L110'] (odstrani oznako verige)."""
    if not isinstance(residues_str, str) or not residues_str.strip():
        return []

    ocisceni = []
    for d in (d.strip() for d in residues_str.split(",")):
        clean = re.sub(r"^[A-Za-z0-9]+:", "", d)
        if clean:
            ocisceni.append(clean)
    return ocisceni


def generiraj_tabelo_ostankov(ref_row, df_prim):
    """Tabela: vrstica = ostanek referenčnega epitopa, stolpci = primerjalni proteini."""
    ref_ostanki = ref_row["Clean_Residues"]
    proteini_prim = df_prim["Protein"].unique()
    ime_ref_stolpca = f"{ref_row['Protein']} - {ref_row['ID Epitopa']}"

    vrstice = []
    for ostanek in ref_ostanki:
        vrstica = {ime_ref_stolpca: ostanek}

        for prot in proteini_prim:
            df_prot = df_prim[df_prim["Protein"] == prot]
            zadetki = [
                str(prim_row["ID Epitopa"])
                for _, prim_row in df_prot.iterrows()
                if ostanek in prim_row["Clean_Residues"]
            ]
            vrstica[prot] = f"{', '.join(zadetki)} = {len(zadetki)}" if zadetki else "/"

        vrstice.append(vrstica)

    return pd.DataFrame(vrstice)


# ==========================================
# 2. SHRANJEVANJE V OBLIKOVAN EXCEL
# ==========================================

def shrani_v_excel(slovar_tabel, ime_datoteke):
    """Vsak DataFrame iz slovarja zapiše v svoj delovni list (ključ = ime lista)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_ostanek = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
    font_data = Font(name="Segoe UI", size=11, color="334155")

    fill_header_ref = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_header_target = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    fill_ostanek = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_match = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    border_thin = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    align_center = Alignment(horizontal="center", vertical="center")

    for ime_sheeta, df_tab in slovar_tabel.items():
        ws = wb.create_sheet(title=str(ime_sheeta)[:31])

        headers = list(df_tab.columns)
        ws.append(headers)
        for _, row in df_tab.iterrows():
            ws.append(list(row))

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_cell
            cell.fill = fill_header_ref if col_num == 1 else fill_header_target

        # Celice z zadetkom (ne "/") so obarvane zeleno
        for row_num in range(2, len(df_tab) + 2):
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.border = border_cell
                cell.alignment = align_center
                val = str(cell.value or "")

                if col_num == 1:
                    cell.font = font_ostanek
                    cell.fill = fill_ostanek
                else:
                    cell.font = font_data
                    if val not in ("/", ""):
                        cell.fill = fill_match
                    elif row_num % 2 == 0:
                        cell.fill = fill_zebra

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 5, 14)

        ws.freeze_panes = "B2"

    wb.save(ime_datoteke)
    print(f"[OK] Shranjeno v: {ime_datoteke}")


# ==========================================
# 3. GLAVNA IZVEDBA
# ==========================================

def dodaj_koncnico(ime, koncnica):
    return ime if ime.endswith(koncnica) else ime + koncnica


def ime_lista(row):
    """Ime delovnega lista brez znakov, ki jih Excel ne dovoli, in največ 31 znakov."""
    ime = f"{row['Protein']}_{row['ID Epitopa']}"
    for znak in PREPOVEDANI_ZNAKI_SHEET:
        ime = ime.replace(znak, "_")
    return ime[:31]


def main():
    if len(sys.argv) < 4:
        print("Uporaba: python konformacijski_epitopi.py <referenca.csv> <primerjalni.csv> <izhod.xlsx>")
        sys.exit(1)

    pot_ref = dodaj_koncnico(sys.argv[1], ".csv")
    pot_prim = dodaj_koncnico(sys.argv[2], ".csv")
    pot_izhod = dodaj_koncnico(sys.argv[3], ".xlsx")

    for pot in (pot_ref, pot_prim):
        if not os.path.exists(pot):
            print(f"Napaka: datoteka '{pot}' ne obstaja.")
            sys.exit(1)

    print(f"Berem referenčno datoteko: {pot_ref}")
    df_ref = preberi_csv_varno(pot_ref)
    print(f"Berem primerjalno datoteko: {pot_prim}")
    df_prim = preberi_csv_varno(pot_prim)

    for ime, df in [("referenčni", df_ref), ("primerjalni", df_prim)]:
        manjkajoci = [s for s in ZAHTEVANI_STOLPCI if s not in df.columns]
        if manjkajoci:
            print(f"Napaka: v {ime} datoteki manjkajo stolpci: {manjkajoci}")
            print(f"Najdeni stolpci: {df.columns.tolist()}")
            sys.exit(1)

    df_ref["Clean_Residues"] = df_ref["Residues"].apply(ocisti_residues)
    df_prim["Clean_Residues"] = df_prim["Residues"].apply(ocisti_residues)

    print("Izvajam primerjavo ostankov...")
    slovar_vseh_tabel = {}
    for _, row in df_ref.iterrows():
        ime = ime_lista(row)
        if ime in slovar_vseh_tabel:
            print(f"OPOZORILO: ime lista '{ime}' se podvaja, prejšnji podatki bodo prepisani.")
        slovar_vseh_tabel[ime] = generiraj_tabelo_ostankov(row, df_prim)

    shrani_v_excel(slovar_vseh_tabel, pot_izhod)


if __name__ == "__main__":
    main()
