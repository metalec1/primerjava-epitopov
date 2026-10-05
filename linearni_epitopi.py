"""
Primerjava linearnih epitopov dveh proteinov.

Vsak epitop se razčleni na prekrivajoče se dvočrkovne kombinacije (dimere),
nato se za vsak epitop referenčnega proteina preveri, kolikokrat se njegovi
dimeri pojavijo v epitopih primerjalnega proteina. Rezultat je Excel datoteka
z enim delovnim listom na referenčni epitop.

Uporaba:
    python linearni_epitopi.py <referenca.csv> <primerjalni.csv> <izhod.xlsx>
"""

import os
import sys
from collections import Counter

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ZAHTEVANI_STOLPCI = ["Protein", "ID Epitopa", "Sekvenca"]
PREPOVEDANI_ZNAKI_SHEET = [":", "/", "\\", "?", "*", "[", "]"]


# ==========================================
# 1. OBDELAVA SEKVENC
# ==========================================

def prestej_pare(sekvenca):
    """Vrne slovar {dimer: število pojavitev} za prekrivajoče se dimere v sekvenci."""
    pari = [sekvenca[i:i + 2] for i in range(len(sekvenca) - 1)]
    return dict(Counter(pari))


def generiraj_primerjalno_tabelo(ref_row, df_target):
    """Tabela: vrstica = dimer referenčnega epitopa, stolpci = pojavitve v vsakem primerjalnem epitopu."""
    ref_label = f"{ref_row['Protein']} - {ref_row['ID Epitopa']}"
    ref_dict = ref_row["Pojavitve"]

    podatki = []
    for par, count in ref_dict.items():
        vrstica = {"Par": par, ref_label: count}

        for _, target_row in df_target.iterrows():
            target_label = f"{target_row['Protein']} - {target_row['ID Epitopa']}"
            vrstica[target_label] = target_row["Pojavitve"].get(par, 0)

        podatki.append(vrstica)

    return pd.DataFrame(podatki)


# ==========================================
# 2. SHRANJEVANJE V OBLIKOVAN EXCEL
# ==========================================

def shrani_v_excel(slovar_tabel, ime_datoteke):
    """Vsak DataFrame iz slovarja zapiše v svoj delovni list (ključ = ime lista)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_par = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
    font_data = Font(name="Segoe UI", size=11, color="334155")

    fill_header_ref = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_header_target = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    fill_par = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_match = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    border_thin = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    align_center = Alignment(horizontal="center", vertical="center")

    for ime_sheeta, df_tab in slovar_tabel.items():
        ws = wb.create_sheet(title=str(ime_sheeta)[:31])

        headers = list(df_tab.columns)
        ws.append(headers)
        for _, row in df_tab.iterrows():
            ws.append(list(row))

        # Glava: prva dva stolpca (dimer + referenca) temno modra, ostali sivi
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_cell
            cell.fill = fill_header_ref if col_num <= 2 else fill_header_target

        # Podatki: ujemanja (> 0) v primerjalnih stolpcih so obarvana zeleno
        for row_num in range(2, len(df_tab) + 2):
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.border = border_cell
                cell.alignment = align_center
                val = cell.value

                if col_num == 1:
                    cell.font = font_par
                    cell.fill = fill_par
                else:
                    cell.font = font_data
                    if col_num > 2 and isinstance(val, (int, float)) and val > 0:
                        cell.fill = fill_match
                    elif row_num % 2 == 0:
                        cell.fill = fill_zebra

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 12)

        ws.freeze_panes = "C2"

    wb.save(ime_datoteke)
    print(f"[OK] Shranjeno v: {ime_datoteke}")


# ==========================================
# 3. GLAVNA IZVEDBA
# ==========================================

def dodaj_koncnico(ime, koncnica):
    return ime if ime.endswith(koncnica) else ime + koncnica


def ime_lista(row):
    """Ime delovnega lista brez znakov, ki jih Excel ne dovoli, in največ 31 znakov."""
    ime = f"{row['Protein']}_{row['ID Epitopa']}"
    for znak in PREPOVEDANI_ZNAKI_SHEET:
        ime = ime.replace(znak, "_")
    return ime[:31]


def main():
    if len(sys.argv) < 4:
        print("Uporaba: python linearni_epitopi.py <referenca.csv> <primerjalni.csv> <izhod.xlsx>")
        sys.exit(1)

    pot_ref = dodaj_koncnico(sys.argv[1], ".csv")
    pot_prim = dodaj_koncnico(sys.argv[2], ".csv")
    pot_izhod = dodaj_koncnico(sys.argv[3], ".xlsx")

    for pot in (pot_ref, pot_prim):
        if not os.path.exists(pot):
            print(f"Napaka: datoteka '{pot}' ne obstaja.")
            sys.exit(1)

    print(f"Berem referenčno datoteko: {pot_ref}")
    df_ref = pd.read_csv(pot_ref)
    print(f"Berem primerjalno datoteko: {pot_prim}")
    df_prim = pd.read_csv(pot_prim)

    for ime, df in [("referenčni", df_ref), ("primerjalni", df_prim)]:
        manjkajoci = [s for s in ZAHTEVANI_STOLPCI if s not in df.columns]
        if manjkajoci:
            print(f"Napaka: v {ime} datoteki manjkajo stolpci: {manjkajoci}")
            print(f"Najdeni stolpci: {df.columns.tolist()}")
            sys.exit(1)

    df_ref["Pojavitve"] = df_ref["Sekvenca"].apply(prestej_pare)
    df_prim["Pojavitve"] = df_prim["Sekvenca"].apply(prestej_pare)

    print("Izvajam primerjavo sekvenc...")
    slovar_vseh_tabel = {}
    for _, row in df_ref.iterrows():
        ime = ime_lista(row)
        if ime in slovar_vseh_tabel:
            print(f"OPOZORILO: ime lista '{ime}' se podvaja, prejšnji podatki bodo prepisani.")
        slovar_vseh_tabel[ime] = generiraj_primerjalno_tabelo(row, df_prim)

    shrani_v_excel(slovar_vseh_tabel, pot_izhod)


if __name__ == "__main__":
    main()
